from __future__ import annotations

import os
from concurrent.futures import ProcessPoolExecutor, as_completed
from pathlib import Path

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook

SRC = Path(r"C:\Users\Grigosha\Downloads\YourBoardingPassDotAero")
OUT_DIR = Path("../parquet_out")
OUT_DIR.mkdir(exist_ok=True)


def val(ws, cell: str):
    v = ws[cell].value
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def parse_date(s):
    if not s:
        return None
    try:
        return pd.to_datetime(s).date()
    except Exception:
        return None


def parse_time(s):
    if not s:
        return None
    try:
        return pd.to_datetime(s, format="%H:%M").time()
    except Exception:
        try:
            return pd.to_datetime(s).time()
        except Exception:
            return None


def parse_sheet(ws):
    if val(ws, "A1") != "BOARDING PASS":
        return None

    pnr = et = None
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.strip() == "PNR:":
                pnr = val(ws, f"{chr(64 + c.column + 1)}{c.row}")
                et = val(ws, f"{chr(64 + c.column + 4)}{c.row}")

    seq_raw = val(ws, "H1")
    try:
        seq = int(seq_raw) if seq_raw is not None else None
    except ValueError:
        seq = None

    return {
        "sequence": seq,
        "title": val(ws, "A3"),
        "passenger_name": val(ws, "B3"),
        "ff_number": val(ws, "F3"),
        "class": val(ws, "H3"),
        "flight": val(ws, "A5"),
        "from_city": val(ws, "D5"),
        "to_city": val(ws, "H5"),
        "gate": val(ws, "B7"),
        "from_iata": val(ws, "D7"),
        "to_iata": val(ws, "H7"),
        "flight_date": parse_date(val(ws, "A9")),
        "departure_time": parse_time(val(ws, "C9")),
        "operator": val(ws, "F9"),
        "seat": val(ws, "H11"),
        "pnr": pnr,
        "e_ticket": et,
    }


def process_file(path_str: str) -> tuple[str, list[dict], str | None]:
    path = Path(path_str)
    rows: list[dict] = []
    err = None
    try:
        wb = load_workbook(path, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            try:
                rec = parse_sheet(ws)
            except Exception as e:
                err = f"{sheet_name}: {e}"
                continue
            if rec is None:
                continue
            rec["source_file"] = path.name
            rec["source_sheet"] = sheet_name
            rows.append(rec)
        wb.close()
    except Exception as e:
        err = f"open: {e}"
    return path.name, rows, err


SCHEMA = pa.schema([
    ("sequence", pa.int64()),
    ("title", pa.string()),
    ("passenger_name", pa.string()),
    ("ff_number", pa.string()),
    ("class", pa.string()),
    ("flight", pa.string()),
    ("from_city", pa.string()),
    ("to_city", pa.string()),
    ("gate", pa.string()),
    ("from_iata", pa.string()),
    ("to_iata", pa.string()),
    ("flight_date", pa.date32()),
    ("departure_time", pa.string()),
    ("operator", pa.string()),
    ("seat", pa.string()),
    ("pnr", pa.string()),
    ("e_ticket", pa.string()),
    ("source_file", pa.string()),
    ("source_sheet", pa.string()),
])


def rows_to_table(rows: list[dict]) -> pa.Table:
    norm = []
    for r in rows:
        rr = dict(r)
        rr["departure_time"] = r["departure_time"].strftime("%H:%M") if r["departure_time"] else None
        norm.append(rr)
    return pa.Table.from_pylist(norm, schema=SCHEMA)


def main():
    files = sorted(SRC.rglob("*.xlsx"))
    print(f"Found {len(files)} files")

    workers = max(1, (os.cpu_count() or 4))
    print(f"Workers: {workers}")

    writer: pq.ParquetWriter | None = None
    total_rows = 0
    errors: list[str] = []

    with ProcessPoolExecutor(max_workers=workers) as ex:
        futures = {ex.submit(process_file, str(f)): f for f in files}
        for i, fut in enumerate(as_completed(futures), 1):
            name, rows, err = fut.result()
            if err:
                errors.append(f"{name}: {err}")
            if rows:
                table = rows_to_table(rows)
                if writer is None:
                    writer = pq.ParquetWriter(OUT_DIR / "boarding_passes.parquet", SCHEMA)
                writer.write_table(table)
                total_rows += len(rows)
            if i % 25 == 0 or i == len(files):
                print(f"  {i}/{len(files)} files, {total_rows} rows")

    if writer:
        writer.close()

    print(f"\nDone. Rows: {total_rows}")
    if errors:
        print(f"Errors ({len(errors)}):")
        for e in errors[:20]:
            print(" ", e)


if __name__ == "__main__":
    main()